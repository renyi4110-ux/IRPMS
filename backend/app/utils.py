from __future__ import annotations

import io
import json
import csv
from datetime import datetime
from difflib import SequenceMatcher
from typing import Dict, List, Optional

import pandas as pd
from docx import Document
from openpyxl import load_workbook

ISSUE_KEYWORDS = ["失败", "出错", "不足", "无法", "问题", "风险", "卡住", "延迟"]
COMPLETE_KEYWORDS = ["完成", "已完成", "结束"]
PROGRESS_KEYWORDS = ["推进", "进行", "开展", "完成了", "实现"]


def load_text_from_file(file_bytes: bytes, filename: str) -> str:
    lower_name = filename.lower()
    if lower_name.endswith((".txt", ".md")):
        return file_bytes.decode("utf-8", errors="ignore")
    if lower_name.endswith(".docx"):
        doc = Document(io.BytesIO(file_bytes))
        return "\n".join([p.text for p in doc.paragraphs])
    raise ValueError("Unsupported text format")


def parse_task_plan(file_bytes: bytes, filename: str) -> List[Dict[str, str]]:
    lower_name = filename.lower()
    records: List[Dict[str, str]] = []
    if lower_name.endswith(".csv"):
        csv_file = io.StringIO(file_bytes.decode("utf-8", errors="ignore"))
        reader = csv.DictReader(csv_file)
        for row in reader:
            records.append(row)
    elif lower_name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(file_bytes))
        ws = wb.active
        headers = [str(cell.value).strip() if cell.value else "" for cell in next(ws.iter_rows(max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            record = {headers[i]: row[i] for i in range(len(headers))}
            records.append(record)
    else:
        raise ValueError("Unsupported plan format")
    return records


def parse_text_records(text: str) -> List[str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return lines


def get_similarity(a: str, b: str) -> float:
    return SequenceMatcher(None, a, b).ratio()


def match_task_name(candidate: str, tasks: List[str]) -> Optional[str]:
    candidate = candidate.strip()
    if not candidate:
        return None
    best_name = None
    best_score = 0
    for task_name in tasks:
        score = get_similarity(candidate, task_name)
        if score > best_score:
            best_name = task_name
            best_score = score
    if best_score >= 0.55:
        return best_name
    return None


def detect_issue(text: str) -> Optional[str]:
    for keyword in ISSUE_KEYWORDS:
        if keyword in text:
            return text
    return None


def detect_status(text: str) -> str:
    for keyword in COMPLETE_KEYWORDS:
        if keyword in text:
            return "已完成"
    for keyword in PROGRESS_KEYWORDS:
        if keyword in text:
            return "进行中"
    return "进行中"


def guess_owner(text: str, fallback: Optional[str] = None) -> str:
    separators = ["负责人", "执行", "由", "学生", "成员"]
    for sep in separators:
        if sep in text:
            idx = text.index(sep)
            snippet = text[idx : idx + 15]
            for symbol in [":", "："]:
                if symbol in snippet:
                    owner = snippet.split(symbol)[-1]
                    return owner.strip()[:5] or fallback or "未指定"
    return fallback or "未指定"


def build_summary(task_logs: List[str]) -> str:
    if not task_logs:
        return ""
    if len(task_logs) == 1:
        return task_logs[0][:150]
    return "；".join(entry[:80] for entry in task_logs[-3:])


def analyze_records(records: List[str], known_tasks: List[str]) -> List[Dict[str, str]]:
    payloads: List[Dict[str, str]] = []
    for text in records:
        task_name = None
        for marker in ["任务", "目标", "事项"]:
            if marker in text:
                task_name = text.split(marker, 1)[-1].strip()
                break
        if not task_name:
            task_name = text[:20]
        matched_name = match_task_name(task_name, known_tasks)
        final_name = matched_name or task_name
        issue_text = detect_issue(text)
        status = detect_status(text)
        owner = guess_owner(text)
        payloads.append(
            {
                "task_name": final_name,
                "owner": owner,
                "status": status,
                "summary": text[:200],
                "has_issue": bool(issue_text),
                "issue_text": issue_text or "",
            }
        )
    return payloads


def plan_record_to_task(record: Dict[str, str]) -> Dict[str, Optional[str]]:
    name = record.get("任务名称") or record.get("任务") or record.get("name") or record.get("task")
    owner = record.get("负责人") or record.get("owner") or record.get("负责人/执行人")
    start = record.get("计划开始时间") or record.get("plan_start")
    end = record.get("计划结束时间") or record.get("plan_end")
    def _parse_date(value: Optional[str]):
        if value in (None, ""):
            return None
        if isinstance(value, datetime):
            return value.date()
        try:
            return datetime.fromisoformat(str(value)).date()
        except ValueError:
            try:
                return datetime.strptime(str(value), "%Y-%m-%d").date()
            except ValueError:
                return None
    return {
        "name": (name or "未命名任务").strip(),
        "owner": (owner or "未指定").strip(),
        "plan_start": _parse_date(start),
        "plan_end": _parse_date(end),
    }


def create_notification_message(kind: str, payload: Dict[str, str]) -> str:
    if kind == "log":
        return f"已解析 {payload.get('records', 0)} 条记录，更新 {payload.get('tasks', 0)} 个任务。"
    if kind == "report":
        return f"{payload.get('title', '报告')} 已生成，点击查看。"
    if kind == "risk":
        return f"发现 {payload.get('risk_tasks', 0)} 个存在风险的任务，请及时关注。"
    return json.dumps(payload, ensure_ascii=False)
